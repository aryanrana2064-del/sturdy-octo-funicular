"""Deposit date / total deposit / deposit balance / SHREEJI BOXES branding.

Every figure checked here is read back from a real SQLite file and a real exported PDF/XLSX.
"""
import re
import shutil
import sqlite3
import subprocess
import tempfile
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from tests.helpers import ServiceTestCase
from vx7khata import backup, db, exports, money
from vx7khata.service import KhataService, ValidationError

try:
    from pypdf import PdfReader
    HAVE_PYPDF = True
except ImportError:  # pragma: no cover
    HAVE_PYPDF = False

GENERATED = datetime(2026, 10, 3, 14, 30)


def paise(cell_value) -> int:
    return round(float(cell_value) * 100)


def pdf_text(path) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(str(path)).pages)


def watermark_per_page(path) -> list[bool]:
    """True for each page that carries the big SJB watermark (text of 150pt or more)."""
    result = []
    for page in PdfReader(str(path)).pages:
        found = []

        def visitor(text, cm, tm, font_dict, font_size, _found=found):
            if text.strip() == "SJB" and font_size >= 150:
                _found.append(font_size)

        page.extract_text(visitor_text=visitor)
        result.append(bool(found))
    return result


class DepositTestCase(ServiceTestCase):
    def setUp(self):
        super().setUp()
        self.c = self.svc.add_customer("Rahul Sharma", "9876543210", "12 Main Road")
        self.amit = self.svc.add_customer("Amit")
        self.out = Path(self.tmp.name)

    def history(self):
        """Udhaar Jan/Mar, and a deposit that was ENTERED LAST but really happened on 20-01-2025."""
        self.u1 = self.add(self.c.id, "UDHAAR", "10-01-2025", "Gold box set", "5000")
        self.u2 = self.add(self.c.id, "UDHAAR", "10-03-2025", "Cartons", "3000")
        self.d = self.add(self.c.id, "JAMA", "01-06-2025", "Cash received", "2000",
                          deposit_date="20-01-2025", notes="old payment")


# =============================================================================== 2. deposit date
class TestDepositDate(DepositTestCase):
    def test_deposit_date_is_stored_and_separate_from_dates_and_creation_time(self):
        self.history()
        t = self.svc.get_transaction(self.d.id)
        self.assertEqual(t.deposit_date, "2025-01-20")
        self.assertEqual(t.txn_date, "2025-06-01")  # untouched by the deposit date
        self.assertTrue(t.created_at.startswith("2026-10-03"))  # real entry time, not the deposit date
        raw = self.svc.conn.execute("SELECT deposit_date FROM transactions WHERE id=?", (self.d.id,)).fetchone()
        self.assertEqual(raw[0], "2025-01-20")  # persisted in SQLite

    def test_udhaar_rows_have_no_deposit_date(self):
        self.history()
        self.assertIsNone(self.svc.get_transaction(self.u1.id).deposit_date)
        self.assertIsNone(self.svc.ledger(self.c.id).rows[0].deposit_date)

    def test_blank_deposit_date_defaults_to_entry_date(self):
        t = self.add(self.c.id, "JAMA", "05-02-2025", "Payment", "100")
        self.assertEqual(t.deposit_date, "2025-02-05")

    def test_previous_years_are_accepted(self):
        t = self.add(self.c.id, "JAMA", "05-02-2025", "Payment", "100", deposit_date="01-04-2019")
        self.assertEqual(t.deposit_date, "2019-04-01")

    def test_edit_deposit_date_keeps_created_at_and_entry_date(self):
        self.history()
        before = self.svc.get_transaction(self.d.id)
        self.clock.advance(3600)
        after = self.svc.update_transaction(self.d.id, "JAMA", "01-06-2025", "Cash received", amount="2000",
                                            notes="old payment", deposit_date="15-03-2025")
        self.assertEqual(after.deposit_date, "2025-03-15")
        self.assertEqual(after.txn_date, before.txn_date)
        self.assertEqual(after.created_at, before.created_at)
        self.assertNotEqual(after.updated_at, before.updated_at)

    def test_validation(self):
        bad = ("31-02-2025", "04-10-2026", "10-01-1800", "tomorrow", "2025-13-45")
        for value in bad:
            with self.subTest(value):
                with self.assertRaises(ValidationError):
                    self.add(self.c.id, "JAMA", "05-02-2025", "x", "100", deposit_date=value)
        with self.assertRaises(ValidationError):  # deposit date makes no sense on an udhaar
            self.add(self.c.id, "UDHAAR", "05-02-2025", "x", "100", deposit_date="01-01-2025")
        self.assertEqual(self.svc.ledger(self.c.id).rows, [])  # nothing half-saved
        with self.assertRaises(ValidationError):  # amount validation still applies to deposits
            self.add(self.c.id, "JAMA", "05-02-2025", "x", "-5", deposit_date="01-01-2025")

    def test_same_payment_on_another_deposit_date_is_not_a_duplicate(self):
        self.svc.add_transaction(self.c.id, "JAMA", "05-02-2025", "Cash", amount="100", deposit_date="01-02-2025")
        self.clock.advance(2)
        self.svc.add_transaction(self.c.id, "JAMA", "05-02-2025", "Cash", amount="100", deposit_date="02-02-2025")
        self.assertEqual(len(self.svc.ledger(self.c.id).rows), 2)

    def test_legacy_deposit_without_a_date_is_handled_safely(self):
        self.add(self.c.id, "UDHAAR", "10-01-2025", "A", "1000")
        self.svc.conn.execute(
            "INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at) "
            "VALUES (?, '2025-02-10', 'Legacy payment', 40000, 'JAMA', 'n', 'n')", (self.c.id,))
        led = self.svc.ledger(self.c.id)
        self.assertEqual([r.txn.item for r in led.rows], ["A", "Legacy payment"])
        self.assertEqual(led.rows[1].deposit_date, "2025-02-10")  # falls back to the entry date
        self.assertEqual([r.balance for r in led.rows], [100000, 60000])
        pdf = exports.export_statement_pdf(led, self.out / "l.pdf", generated_at=GENERATED)
        xl = exports.export_statement_xlsx(led, self.out / "l.xlsx", generated_at=GENERATED)
        self.assertTrue(pdf.exists() and xl.exists())
        ws = load_workbook(xl, data_only=True)["Statement"]
        legacy = next(r for r in ws.iter_rows(values_only=True) if str(r[1]).startswith("Legacy payment"))
        self.assertEqual(legacy[4].date(), date(2025, 2, 10))


# =============================================================================== 4. chronological balance
class TestChronologicalBalance(DepositTestCase):
    def test_old_deposit_lands_at_its_true_place(self):
        self.history()
        led = self.svc.ledger(self.c.id)
        self.assertEqual([r.txn.item for r in led.rows], ["Gold box set", "Cash received", "Cartons"])
        self.assertEqual([r.balance for r in led.rows], [500000, 300000, 600000])
        self.assertEqual(led.closing_balance, 600000)

    def test_outstanding_equals_total_udhaar_minus_total_deposit(self):
        self.history()
        led = self.svc.ledger(self.c.id)
        self.assertEqual(led.total_udhaar, 800000)
        self.assertEqual(led.total_deposit, 200000)
        self.assertEqual(led.deposit_balance, led.opening_balance + led.total_udhaar - led.total_deposit)
        self.assertEqual(self.svc.customer_balance(self.c.id), 600000)

    def test_editing_an_old_deposit_recalculates(self):
        self.history()
        self.svc.update_transaction(self.d.id, "JAMA", "01-06-2025", "Cash received", amount="2000",
                                    notes="old payment", deposit_date="15-03-2025")
        led = self.svc.ledger(self.c.id)
        self.assertEqual([r.txn.item for r in led.rows], ["Gold box set", "Cartons", "Cash received"])
        self.assertEqual([r.balance for r in led.rows], [500000, 800000, 600000])
        # amount change flows through too
        self.svc.update_transaction(self.d.id, "JAMA", "01-06-2025", "Cash received", amount="3500",
                                    deposit_date="15-03-2025")
        self.assertEqual([r.balance for r in self.svc.ledger(self.c.id).rows], [500000, 800000, 450000])

    def test_deleting_an_old_deposit_recalculates(self):
        self.history()
        self.svc.delete_transaction(self.d.id)
        self.assertEqual([r.balance for r in self.svc.ledger(self.c.id).rows], [500000, 800000])

    def test_adding_an_old_deposit_later_shifts_every_later_balance(self):
        self.history()
        self.add(self.c.id, "JAMA", "02-06-2025", "Another old payment", "1000", deposit_date="01-02-2025")
        led = self.svc.ledger(self.c.id)
        self.assertEqual([r.txn.item for r in led.rows],
                         ["Gold box set", "Cash received", "Another old payment", "Cartons"])
        self.assertEqual([r.balance for r in led.rows], [500000, 300000, 200000, 500000])

    def test_never_mixes_customers(self):
        self.history()
        self.add(self.amit.id, "UDHAAR", "01-02-2025", "Amit box", "999")
        self.add(self.amit.id, "JAMA", "01-06-2025", "Amit pays", "100", deposit_date="05-02-2025")
        rahul = self.svc.ledger(self.c.id)
        amit = self.svc.ledger(self.amit.id)
        self.assertEqual([r.balance for r in rahul.rows], [500000, 300000, 600000])
        self.assertEqual([r.balance for r in amit.rows], [99900, 89900])
        self.assertEqual((amit.total_udhaar, amit.total_deposit), (99900, 10000))
        self.assertEqual(self.svc.customer_balance(self.amit.id), 89900)

    def test_date_range_uses_deposit_date_and_correct_opening_balance(self):
        self.history()
        led = self.svc.ledger(self.c.id, date_from="15-01-2025", date_to="31-01-2025")
        self.assertEqual([r.txn.item for r in led.rows], ["Cash received"])  # entry date is June, deposit is Jan
        self.assertEqual((led.opening_balance, led.total_udhaar, led.total_deposit, led.closing_balance),
                         (500000, 0, 200000, 300000))
        led = self.svc.ledger(self.c.id, date_from="01-02-2025")
        self.assertEqual((led.opening_balance, led.total_deposit), (300000, 0))
        self.assertEqual([r.txn.item for r in led.rows], ["Cartons"])
        led = self.svc.ledger(self.c.id, date_from="01-06-2025", date_to="30-06-2025")
        self.assertEqual(led.rows, [])  # the June *entry date* does not count: the money arrived in January

    def test_total_deposit_sums_only_deposits_in_the_range(self):
        self.history()
        self.add(self.c.id, "JAMA", "02-06-2025", "More", "700", deposit_date="12-03-2025")
        led = self.svc.ledger(self.c.id, date_from="01-01-2025", date_to="31-03-2025")
        self.assertEqual(led.total_deposit, 270000)
        self.assertEqual(led.total_udhaar, 800000)
        reg = self.svc.register(date_from="01-03-2025", date_to="31-03-2025", customer_ids=[self.c.id])
        self.assertEqual((reg.total_udhaar, reg.total_deposit), (300000, 70000))

    def test_dashboard_counts_a_deposit_on_its_deposit_date(self):
        self.add(self.c.id, "JAMA", "03-10-2026", "Paid today for an old bill", "50", deposit_date="03-10-2026")
        self.add(self.c.id, "JAMA", "03-10-2026", "Recorded today, paid long ago", "60", deposit_date="01-01-2025")
        d = self.svc.dashboard()
        self.assertEqual((d.todays_transactions, d.historical_transactions), (1, 1))


# =============================================================================== 7. migration
class TestMigration(ServiceTestCase):
    def make_v2(self, path):
        conn = sqlite3.connect(path, isolation_level=None)
        for v in (1, 2):
            for stmt in db.MIGRATIONS[v]:
                conn.execute(stmt)
        conn.execute("PRAGMA user_version = 2")
        conn.execute("INSERT INTO customers(name, name_key, mobile, created_at, updated_at) "
                     "VALUES('Old Customer', 'old customer', '9876543210', '2024-12-01T10:00:00+05:30', 'n')")
        conn.execute("INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at) "
                     "VALUES (1, '2025-01-10', 'Kept udhaar', 500000, 'UDHAAR', '2025-01-11T10:00:00+05:30', 'n')")
        conn.execute("INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at) "
                     "VALUES (1, '2025-01-20', 'Kept payment', 200000, 'JAMA', '2025-01-21T09:30:00+05:30', 'n')")
        conn.close()

    def test_v2_database_upgrades_keeping_everything(self):
        path = Path(self.tmp.name) / "v2.db"
        self.make_v2(path)
        svc = KhataService(path)
        self.addCleanup(svc.close)
        self.assertEqual(db.get_user_version(svc.conn), db.SCHEMA_VERSION)
        cols = {r[1] for r in svc.conn.execute("PRAGMA table_info(transactions)")}
        self.assertIn("deposit_date", cols)
        led = svc.ledger(1)
        self.assertEqual([(r.txn.item, r.txn.txn_date, r.txn.deposit_date, r.balance) for r in led.rows],
                         [("Kept udhaar", "2025-01-10", None, 500000),
                          ("Kept payment", "2025-01-20", "2025-01-20", 300000)])  # deposit date back-filled
        self.assertEqual(led.rows[1].txn.created_at, "2025-01-21T09:30:00+05:30")  # creation time untouched
        self.assertEqual(svc.customer_balance(1), 300000)
        bak = Path(self.tmp.name) / "v2.db.pre-v2-migration.bak"
        self.assertTrue(bak.exists())
        old = sqlite3.connect(bak)
        self.addCleanup(old.close)
        self.assertNotIn("deposit_date", {r[1] for r in old.execute("PRAGMA table_info(transactions)")})

    def test_migration_is_idempotent_on_reopen(self):
        path = Path(self.tmp.name) / "v2.db"
        self.make_v2(path)
        KhataService(path).close()
        again = KhataService(path)
        self.addCleanup(again.close)
        self.assertEqual(again.ledger(1).closing_balance, 300000)

    def test_deposit_date_check_constraint(self):
        c = self.svc.add_customer("A")
        with self.assertRaises(sqlite3.IntegrityError):
            self.svc.conn.execute(
                "INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at, "
                "deposit_date) VALUES (?, '2025-01-01', 'x', 100, 'JAMA', 'n', 'n', 'not-a-date')", (c.id,))


class TestBackupWithDepositDates(DepositTestCase):
    def test_backup_and_restore_keep_deposit_dates(self):
        self.history()
        bk = self.out / "b.db"
        self.svc.create_backup(bk)
        self.add(self.c.id, "UDHAAR", "01-04-2025", "later", "1")
        self.svc.restore_backup(bk, safety_dir=self.out / "safe")
        led = self.svc.ledger(self.c.id)
        self.assertEqual([r.deposit_date for r in led.rows], [None, "2025-01-20", None])
        self.assertEqual([r.balance for r in led.rows], [500000, 300000, 600000])

    def test_restoring_a_v2_backup_upgrades_it(self):
        old = self.out / "v2.db"
        TestMigration.make_v2(self, old)
        self.svc.restore_backup(old, safety_dir=self.out / "safe")
        led = self.svc.ledger(1)
        self.assertEqual(led.rows[1].txn.deposit_date, "2025-01-20")
        self.assertEqual(led.closing_balance, 300000)

    def test_backup_with_invalid_deposit_date_is_rejected(self):
        self.history()
        bk = self.out / "b.db"
        self.svc.create_backup(bk)
        bad = self.out / "bad.db"
        shutil.copyfile(bk, bad)
        conn = sqlite3.connect(bad)
        conn.execute("PRAGMA ignore_check_constraints = ON")
        conn.execute("UPDATE transactions SET deposit_date = 'garbage' WHERE txn_type = 'JAMA'")
        conn.commit()
        conn.close()
        with self.assertRaises(backup.BackupError):
            backup.validate_backup(bad)


# =============================================================================== 1 + 5. PDF
@unittest.skipUnless(HAVE_PYPDF, "pypdf not installed")
class TestStatementPdf(DepositTestCase):
    def setUp(self):
        super().setUp()
        self.history()
        self.led = self.svc.ledger(self.c.id)
        self.path = exports.export_statement_pdf(self.led, self.out / "s.pdf", generated_at=GENERATED)
        self.text = pdf_text(self.path)

    def test_columns_and_heading_text(self):
        flat = " ".join(self.text.split())
        for heading in ("TRANSACTION DATE", "ITEM / DESCRIPTION", "UDHAAR (₹)", "TOTAL DEPOSIT (₹)", "DEPOSIT DATE",
                        "DEPOSIT BALANCE (₹)", "NOTES"):
            self.assertIn(heading, flat)
        order = [flat.index(h) for h in ("TRANSACTION DATE", "ITEM / DESCRIPTION", "UDHAAR (₹)",
                                         "TOTAL DEPOSIT (₹)", "DEPOSIT DATE", "DEPOSIT BALANCE (₹)", "NOTES")]
        self.assertEqual(order, sorted(order))

    def test_report_contents_come_from_the_saved_ledger(self):
        for needle in ("Rahul Sharma", "9876543210", "12 Main Road", "03-10-2026 14:30", "All dates",
                       "10-01-2025", "20-01-2025", "10-03-2025", "old payment", "Cash received",
                       "5,000.00", "2,000.00", "3,000.00", "8,000.00", "6,000.00", "3,000.00"):
            self.assertIn(needle, self.text, needle)
        for row in self.led.rows:  # every running balance in the DB appears in the PDF
            self.assertIn(money.format_inr(row.balance, symbol=False), self.text)
        self.assertRegex(self.text, r"Total Udhaar\s*\n₹8,000\.00")
        self.assertRegex(self.text, r"Total Deposit\s*\n₹2,000\.00")
        self.assertRegex(self.text, r"Final Deposit Balance\s*\n₹6,000\.00")

    def test_rows_are_in_chronological_order_by_deposit_date(self):
        i1, i2, i3 = (self.text.index(s) for s in ("Gold box set", "Cash received", "Cartons"))
        self.assertTrue(i1 < i2 < i3)

    def test_brand_and_watermark_on_the_page(self):
        self.assertIn("SHREEJI BOXES", self.text)
        self.assertEqual(watermark_per_page(self.path), [True])

    def test_watermark_is_drawn_behind_the_table(self):
        order = []
        page = PdfReader(str(self.path)).pages[0]
        page.extract_text(visitor_text=lambda text, cm, tm, fd, size: order.append((text.strip(), size)))
        first_big = next(i for i, (t, s) in enumerate(order) if t == "SJB" and s >= 150)
        first_table = next(i for i, (t, s) in enumerate(order) if t == "TOTAL")
        self.assertLess(first_big, first_table)

    def test_no_website_unless_configured(self):
        self.assertNotRegex(self.text.lower(), r"https?://|www\.|\.com")
        path = exports.export_statement_pdf(self.led, self.out / "w.pdf", website="www.example.test",
                                            generated_at=GENERATED)
        self.assertIn("www.example.test", pdf_text(path))

    def test_period_filter_and_opening_balance(self):
        led = self.svc.ledger(self.c.id, date_from="01-02-2025", date_to="31-12-2025")
        text = pdf_text(exports.export_statement_pdf(led, self.out / "p.pdf", generated_at=GENERATED))
        self.assertIn("01-02-2025 to 31-12-2025", text)
        self.assertRegex(text, r"Opening Balance\s*\n₹3,000\.00")  # 5,000 udhaar - 2,000 January deposit
        self.assertRegex(text, r"Final Deposit Balance\s*\n₹6,000\.00")

    def test_many_pages_repeat_watermark_header_footer_and_numbering(self):
        for i in range(150):
            self.add(self.c.id, "UDHAAR", "01-04-2025", f"bulk item {i}", "10")
        led = self.svc.ledger(self.c.id)
        path = exports.export_statement_pdf(led, self.out / "big.pdf", generated_at=GENERATED)
        reader = PdfReader(str(path))
        pages = len(reader.pages)
        self.assertGreater(pages, 2)
        self.assertEqual(watermark_per_page(path), [True] * pages)
        for n, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            self.assertIn("SHREEJI BOXES", text)
            self.assertIn(f"Page {n} of {pages}", text)
            if "bulk item" in text:  # every page that carries table rows repeats the column headings
                self.assertIn("TOTAL DEPOSIT (₹)", " ".join(text.split()))

    def test_register_and_customer_reports_are_branded_too(self):
        reg = exports.export_register_pdf(self.svc.register(), self.out / "r.pdf", generated_at=GENERATED)
        cus = exports.export_customers_pdf(self.svc.list_customers(), self.out / "c.pdf", generated_at=GENERATED)
        for path in (reg, cus):
            self.assertIn("SHREEJI BOXES", pdf_text(path))
            self.assertEqual(watermark_per_page(path), [True])
        self.assertIn("20-01-2025", pdf_text(reg))
        flat = " ".join(pdf_text(cus).split())
        self.assertIn("TOTAL DEPOSIT (₹)", flat)
        self.assertIn("DEPOSIT BALANCE (₹)", flat)


# =============================================================================== 1 + 6. Excel
class TestStatementExcel(DepositTestCase):
    def setUp(self):
        super().setUp()
        self.history()
        self.led = self.svc.ledger(self.c.id)
        self.path = exports.export_statement_xlsx(self.led, self.out / "s.xlsx", generated_at=GENERATED)
        self.ws = load_workbook(self.path)["Statement"]  # formulas
        self.wv = load_workbook(self.path, data_only=True)["Statement"]  # cached values

    def header_row(self):
        return next(r for r in range(1, self.ws.max_row + 1) if self.ws.cell(r, 1).value == "TRANSACTION DATE")

    def test_is_a_genuine_xlsx_with_exact_headers(self):
        self.assertTrue(zipfile.is_zipfile(self.path))
        with zipfile.ZipFile(self.path) as z:
            self.assertIsNone(z.testzip())
            self.assertIn("xl/workbook.xml", z.namelist())
        h = self.header_row()
        headers = [self.ws.cell(h, c).value for c in range(1, 8)]
        self.assertEqual(headers, ["TRANSACTION DATE", "ITEM / DESCRIPTION", "UDHAAR (₹)", "TOTAL DEPOSIT (₹)",
                                   "DEPOSIT DATE", "DEPOSIT BALANCE (₹)", "NOTES"])

    def test_values_match_the_saved_ledger_row_for_row(self):
        h = self.header_row()
        rows = [[self.wv.cell(r, c).value for c in range(1, 8)] for r in range(h + 2, h + 2 + len(self.led.rows))]
        for xl, row in zip(rows, self.led.rows):
            t = row.txn
            self.assertEqual(xl[0].date(), date.fromisoformat(t.txn_date))
            self.assertTrue(xl[1].startswith(t.item))
            self.assertEqual(xl[2] if xl[2] is None else paise(xl[2]), row.udhaar or None)
            self.assertEqual(xl[3] if xl[3] is None else paise(xl[3]), row.jama or None)
            self.assertEqual(xl[4].date() if xl[4] else None, date.fromisoformat(row.deposit_date) if row.deposit_date else None)
            self.assertEqual(paise(xl[5]), row.balance)
            self.assertEqual(xl[6] or "", t.notes)

    def test_deposit_column_holds_only_the_deposit_amount(self):
        h = self.header_row()
        deposit_row = h + 3  # opening row, udhaar, then the deposit
        self.assertEqual(paise(self.wv.cell(deposit_row, 4).value), 200000)  # amount only
        self.assertEqual(paise(self.wv.cell(deposit_row, 6).value), 300000)  # balance lives in its own column
        self.assertIsNone(self.wv.cell(h + 2, 4).value)  # an udhaar row has no deposit

    def test_formulas_are_live_and_totals_agree(self):
        h = self.header_row()
        first, last = h + 2, h + 1 + len(self.led.rows)
        self.assertEqual(self.ws.cell(first, 6).value, f"=F{first - 1}+C{first}-D{first}")
        totals = next(r for r in range(h, self.ws.max_row + 1) if self.ws.cell(r, 2).value == "TOTAL")
        self.assertEqual(self.ws.cell(totals, 3).value, f"=SUM(C{first}:C{last})")
        self.assertEqual(self.ws.cell(totals, 4).value, f"=SUM(D{first}:D{last})")
        self.assertEqual([paise(self.wv.cell(totals, c).value) for c in (3, 4, 6)], [800000, 200000, 600000])

    def test_separate_summary_section(self):
        labels = {self.ws.cell(r, 2).value: paise(self.wv.cell(r, 3).value)
                  for r in range(1, self.ws.max_row + 1)
                  if self.ws.cell(r, 2).value in ("Opening Balance", "Total Udhaar", "Total Deposit",
                                                  "Final Deposit Balance") and self.wv.cell(r, 3).value is not None}
        self.assertEqual(labels, {"Opening Balance": 0, "Total Udhaar": 800000, "Total Deposit": 200000,
                                  "Final Deposit Balance": 600000})
        self.assertIn("SUMMARY", [self.ws.cell(r, 2).value for r in range(1, self.ws.max_row + 1)])

    def test_formatting_freeze_filter_widths(self):
        h = self.header_row()
        self.assertEqual(self.ws.freeze_panes, f"A{h + 1}")
        self.assertEqual(self.ws.auto_filter.ref, f"A{h}:G{h + 1 + len(self.led.rows)}")
        money_cell = self.ws.cell(h + 2, 3)
        self.assertIn("₹", money_cell.number_format)
        self.assertIn("#,##0.00", money_cell.number_format)
        self.assertEqual(self.ws.cell(h + 2, 1).number_format, "dd-mm-yyyy")
        self.assertEqual(self.ws.cell(h + 3, 5).number_format, "dd-mm-yyyy")
        self.assertTrue(self.ws.cell(h, 1).font.bold)
        self.assertGreaterEqual(self.ws.column_dimensions["B"].width, 30)
        self.assertGreaterEqual(self.ws.column_dimensions["G"].width, 30)
        totals = next(r for r in range(h, self.ws.max_row + 1) if self.ws.cell(r, 2).value == "TOTAL")
        self.assertTrue(self.ws.cell(totals, 6).font.bold)  # final balance highlighted
        self.assertNotEqual(self.ws.cell(totals, 6).fill.fgColor.rgb, "00000000")

    def test_branding_is_everywhere(self):
        self.assertEqual(self.ws["A1"].value, "SHREEJI BOXES")
        self.assertEqual(self.ws["G1"].value, "SJB")
        self.assertIn("SHREEJI BOXES", self.ws.oddFooter.left.text)
        self.assertIn("Page &P of &N", self.ws.oddFooter.right.text)
        self.assertEqual(self.ws.oddHeader.center.text, "&G")  # SJB watermark picture on every printed page
        with zipfile.ZipFile(self.path) as z:
            media = [n for n in z.namelist() if n.startswith("xl/media/")]
            self.assertEqual(len(media), 2)  # printed watermark + on-screen background
            self.assertTrue(all(z.read(n)[:8] == b"\x89PNG\r\n\x1a\n" for n in media))
            self.assertTrue(any("vmlDrawing" in n for n in z.namelist()))  # header picture part
        self.assertEqual(self.ws.page_setup.orientation, "landscape")
        h = self.header_row()
        self.assertEqual(self.ws.print_title_rows, f"${h}:${h}")  # header row repeats on every printed page

    def test_no_website_unless_configured(self):
        strings = [str(c.value) for row in self.ws.iter_rows() for c in row if c.value]
        self.assertFalse(any("http" in s or "www." in s for s in strings))
        footer = "".join(part.text or "" for part in (self.ws.oddFooter.left, self.ws.oddFooter.center, self.ws.oddFooter.right))
        self.assertNotIn("http", footer)
        self.assertNotIn("www.", footer)
        path = exports.export_statement_xlsx(self.led, self.out / "w.xlsx", website="www.example.test",
                                             generated_at=GENERATED)
        self.assertIn("www.example.test", load_workbook(path)["Statement"].oddFooter.left.text)

    def test_item_filter_keeps_true_balances_as_values(self):
        led = self.svc.ledger(self.c.id, search="cartons")
        ws = load_workbook(exports.export_statement_xlsx(led, self.out / "f.xlsx", generated_at=GENERATED))["Statement"]
        h = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "TRANSACTION DATE")
        self.assertEqual(paise(ws.cell(h + 2, 6).value), 600000)  # true account balance, not a formula
        self.assertEqual(led.rows[0].balance, 600000)

    def test_empty_statement_exports(self):
        led = self.svc.ledger(self.amit.id)
        ws = load_workbook(exports.export_statement_xlsx(led, self.out / "e.xlsx", generated_at=GENERATED),
                           data_only=True)["Statement"]
        totals = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "TOTAL")
        self.assertEqual([ws.cell(totals, c).value for c in (3, 4, 6)], [0, 0, 0])
        pdf = exports.export_statement_pdf(led, self.out / "e.pdf", generated_at=GENERATED)
        self.assertTrue(pdf.exists())

    def test_register_and_customer_workbooks_are_branded_and_use_deposit_headings(self):
        reg = exports.export_register_xlsx(self.svc.register(), self.out / "r.xlsx", generated_at=GENERATED)
        cus = exports.export_customers_xlsx(self.svc.list_customers(), self.out / "c.xlsx", generated_at=GENERATED)
        for path in (reg, cus):
            ws = load_workbook(path).active
            self.assertEqual(ws["A1"].value, "SHREEJI BOXES")
            self.assertEqual(ws.oddHeader.center.text, "&G")
            self.assertIn("SHREEJI BOXES", ws.oddFooter.left.text)
            texts = {c.value for row in ws.iter_rows() for c in row}
            self.assertIn("TOTAL DEPOSIT (₹)", texts)
        ws = load_workbook(cus, data_only=True).active
        total = next(r for r in ws.iter_rows(values_only=True) if r[1] == "Overall total")
        self.assertEqual([paise(total[i]) for i in (4, 5, 6)], [800000 + 0, 200000, 600000])

    def test_excel_and_pdf_agree(self):
        text = pdf_text(exports.export_statement_pdf(self.led, self.out / "s.pdf", generated_at=GENERATED))
        pdf_final = re.search(r"Final Deposit Balance\s*\n₹([\d,\.\-]+)", text).group(1)
        h = self.header_row()
        totals = next(r for r in range(h, self.wv.max_row + 1) if self.wv.cell(r, 2).value == "TOTAL")
        self.assertEqual(money.format_inr(paise(self.wv.cell(totals, 6).value), symbol=False), pdf_final)
        pdf_deposit = re.search(r"Total Deposit\s*\n₹([\d,\.\-]+)", text).group(1)
        self.assertEqual(money.format_inr(paise(self.wv.cell(totals, 4).value), symbol=False), pdf_deposit)

    @unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "LibreOffice not installed")
    def test_formulas_recalculate_to_the_same_numbers(self):
        """Strip the cached values, let a real spreadsheet engine calculate every formula, compare."""
        office = shutil.which("soffice") or shutil.which("libreoffice")
        work = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, work, ignore_errors=True)
        stripped = work / "stripped.xlsx"
        wb = load_workbook(self.path)  # openpyxl keeps formulas but drops cached values on save
        wb.save(stripped)
        outdir = work / "out"
        subprocess.run([office, "--headless", f"-env:UserInstallation={(work / 'profile').as_uri()}",
                        "--convert-to", "xlsx", "--outdir", str(outdir), str(stripped)],
                       check=True, capture_output=True, timeout=240)
        recalculated = load_workbook(outdir / "stripped.xlsx", data_only=True)["Statement"]
        checked = 0
        for row in self.ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    got = recalculated[cell.coordinate].value
                    self.assertIsNotNone(got, cell.coordinate)
                    self.assertAlmostEqual(float(got), float(self.wv[cell.coordinate].value), places=2,
                                           msg=cell.coordinate)
                    checked += 1
        self.assertGreaterEqual(checked, 8)


if __name__ == "__main__":
    unittest.main()
