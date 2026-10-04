"""Pending (aging) report, bulk import from Excel/CSV, and WhatsApp helpers."""
import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from tests.helpers import ServiceTestCase
from vx7khata import importer, whatsapp
from vx7khata.service import ValidationError


class TestPendingReport(ServiceTestCase):  # "today" in these tests is 2026-10-03
    def test_oldest_unpaid_udhaar_after_deposits_cover_the_oldest_first(self):
        c = self.svc.add_customer("Rahul")
        self.add(c.id, "UDHAAR", "01-06-2026", "A", "1000")
        self.add(c.id, "UDHAAR", "01-08-2026", "B", "500")
        self.add(c.id, "JAMA", "10-08-2026", "", "1200")  # covers A fully and 200 of B
        rows = self.svc.pending_report()
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual((r.balance, r.oldest_unpaid_date), (30000, "2026-08-01"))
        self.assertEqual(r.days_pending, 63)
        self.assertEqual((r.last_deposit_date, r.last_deposit_paise), ("2026-08-10", 120000))

    def test_paid_up_and_advance_customers_are_not_listed(self):
        a, b = self.svc.add_customer("Paid"), self.svc.add_customer("Advance")
        self.add(a.id, "UDHAAR", "01-06-2026", "A", "100")
        self.add(a.id, "JAMA", "02-06-2026", "", "100")
        self.add(b.id, "JAMA", "02-06-2026", "", "500")
        self.assertEqual(self.svc.pending_report(), [])

    def test_deposit_before_udhaar_is_credit_for_the_next_udhaar(self):
        c = self.svc.add_customer("Early")
        self.add(c.id, "JAMA", "01-05-2026", "", "300")
        self.add(c.id, "UDHAAR", "01-07-2026", "X", "1000")
        r = self.svc.pending_report()[0]
        self.assertEqual((r.balance, r.oldest_unpaid_date), (70000, "2026-07-01"))

    def test_sorted_oldest_first_and_filters(self):
        a, b, c = (self.svc.add_customer(n) for n in ("Anil", "Bhavin", "Chirag"))
        self.add(a.id, "UDHAAR", "01-09-2026", "x", "100")   # 32 days
        self.add(b.id, "UDHAAR", "01-03-2026", "x", "50")    # 216 days
        self.add(c.id, "UDHAAR", "20-09-2026", "x", "900")   # 13 days
        self.assertEqual([r.customer.name for r in self.svc.pending_report()], ["Bhavin", "Anil", "Chirag"])
        self.assertEqual([r.customer.name for r in self.svc.pending_report(min_days=30)], ["Bhavin", "Anil"])
        self.assertEqual([r.customer.name for r in self.svc.pending_report(search="chi")], ["Chirag"])

    def test_back_dated_deposit_changes_pending_since(self):
        c = self.svc.add_customer("Old")
        self.add(c.id, "UDHAAR", "01-01-2026", "A", "100")
        self.add(c.id, "UDHAAR", "01-02-2026", "B", "100")
        self.add(c.id, "JAMA", "01-10-2026", "", "100", deposit_date="05-01-2026")
        self.assertEqual(self.svc.pending_report()[0].oldest_unpaid_date, "2026-02-01")


class TestImportRows(ServiceTestCase):
    def rows(self):
        return [
            {"line": 2, "customer": "Rahul", "mobile": "9876543210", "txn_type": "UDHAAR", "txn_date": "10-01-2025",
             "item": "Box", "amount": "5000"},
            {"line": 3, "customer": "rahul", "txn_type": "JAMA", "txn_date": "20-01-2025", "item": "", "amount": "2000"},
            {"line": 4, "customer": "Amit", "txn_type": "UDHAAR", "txn_date": "05-02-2025", "item": "Tape",
             "quantity": "2", "rate": "50"},
        ]

    def test_dry_run_saves_nothing_and_reports(self):
        rep = self.svc.import_rows(self.rows(), dry_run=True)
        self.assertEqual((rep.total_rows, rep.imported, rep.errors), (3, 0, []))
        self.assertEqual(sorted(rep.new_customers), ["Amit", "Rahul"])
        self.assertEqual((rep.total_udhaar, rep.total_jama), (510000, 200000))
        self.assertEqual(self.svc.list_customers(), [])

    def test_import_creates_customers_matches_names_and_keeps_balances(self):
        rep = self.svc.import_rows(self.rows())
        self.assertEqual(rep.imported, 3)
        by_name = {s.customer.name: s for s in self.svc.list_customers()}
        self.assertEqual(by_name["Rahul"].customer.mobile, "9876543210")
        self.assertEqual(by_name["Rahul"].balance, 300000)  # "rahul" matched "Rahul"
        self.assertEqual(by_name["Amit"].balance, 10000)
        dep = self.svc.ledger(by_name["Rahul"].customer.id).rows[1]
        self.assertEqual(dep.deposit_date, "2025-01-20")  # deposit date defaults to the row date

    def test_one_bad_row_saves_nothing_and_lists_all_problems(self):
        rows = self.rows() + [
            {"line": 9, "customer": "Zed", "txn_type": "UDHAAR", "txn_date": "31-02-2025", "item": "x", "amount": "5"},
            {"line": 11, "customer": "", "txn_type": "UDHAAR", "txn_date": "01-01-2025", "item": "x", "amount": "5"},
        ]
        rep = self.svc.import_rows(rows)
        self.assertEqual([e[0] for e in rep.errors], [9, 11])
        self.assertEqual(rep.imported, 0)
        self.assertEqual(self.svc.list_customers(), [])

    def test_reimport_is_flagged_as_likely_duplicates(self):
        self.svc.import_rows(self.rows())
        rep = self.svc.import_rows(self.rows(), dry_run=True)
        self.assertEqual(rep.likely_duplicates, 3)
        self.assertEqual(rep.new_customers, [])


class TestImporterFiles(ServiceTestCase):
    def xlsx(self, rows, name="in.xlsx"):
        wb = Workbook()
        ws = wb.active
        for r in rows:
            ws.append(r)
        path = Path(self.tmp.name) / name
        wb.save(path)
        return path

    def test_excel_with_dates_numbers_and_hindi_style_headers(self):
        path = self.xlsx([
            ["Old bahi 2024"],  # title row above the headers is fine
            ["Party", "Phone", "Tarikh", "Particulars", "Udhar", "Jama", "Qty", "Rate", "Remarks"],
            ["Rahul", 9876543210, datetime(2025, 1, 10), "Gold box", 5000.5, None, None, None, "p1"],
            [None, None, None, None, None, None, None, None, None],  # blank line is skipped
            ["Rahul", None, datetime(2025, 1, 20), None, None, 2000, None, None, ""],
        ])
        plan = importer.plan_import(self.svc, path)
        self.assertTrue(plan.ok, plan.errors)
        self.assertEqual([r["line"] for r in plan.rows], [3, 5])
        rep = importer.run_import(self.svc, plan)
        self.assertEqual(rep.imported, 2)
        s = self.svc.list_customers()[0]
        self.assertEqual((s.total_udhaar, s.total_jama), (500050, 200000))
        self.assertEqual(s.customer.mobile, "9876543210")

    def test_csv_with_type_and_amount_columns(self):
        path = Path(self.tmp.name) / "in.csv"
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["Customer", "Date", "Item", "Type", "Amount"])
            w.writerow(["Amit", "05/02/2025", "Tape", "Udhaar", "1,250.00"])
            w.writerow(["Amit", "06-02-2025", "", "JAMA", "250"])
        plan = importer.plan_import(self.svc, path)
        self.assertTrue(plan.ok, plan.errors)
        importer.run_import(self.svc, plan)
        self.assertEqual(self.svc.list_customers()[0].balance, 100000)

    def test_problems_block_the_import_and_name_the_row(self):
        path = self.xlsx([
            ["Customer", "Date", "Item", "Udhaar", "Jama"],
            ["Rahul", "10-01-2025", "Box", 100, 50],       # both filled
            ["Rahul", "99-01-2025", "Box", 100, None],     # bad date
        ])
        plan = importer.plan_import(self.svc, path)
        self.assertFalse(plan.ok)
        self.assertEqual([e[0] for e in plan.errors], [2, 3])
        with self.assertRaises(importer.ImportFileError):
            importer.run_import(self.svc, plan)
        self.assertEqual(self.svc.list_customers(), [])

    def test_wrong_file_types_and_missing_columns_are_explained(self):
        bad = Path(self.tmp.name) / "x.xls"
        bad.write_bytes(b"x")
        with self.assertRaises(importer.ImportFileError):
            importer.load_rows(bad)
        with self.assertRaises(importer.ImportFileError):
            importer.load_rows(self.xlsx([["Customer", "Item", "Udhaar"], ["a", "b", 1]], "nodate.xlsx"))

    def test_template_has_the_headers_and_imports_cleanly_when_filled(self):
        path = importer.write_template(Path(self.tmp.name) / "t.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames[0], "Data")
        ws = wb["Data"]
        ws.append(["Rahul", "", "10-01-2025", "Box", 100])
        wb.save(path)
        plan = importer.plan_import(self.svc, path)
        self.assertTrue(plan.ok, plan.errors)


class TestWhatsApp(ServiceTestCase):
    def test_phone_cleaning(self):
        n = whatsapp.normalize_phone
        self.assertEqual(n("98765 43210"), "919876543210")
        self.assertEqual(n("09876543210"), "919876543210")
        self.assertEqual(n("+91 98765-43210"), "919876543210")
        self.assertEqual(n("+1 (415) 555-0100"), "14155550100")
        self.assertIsNone(n(""))
        self.assertIsNone(n("123"))

    def test_link_and_messages(self):
        url = whatsapp.chat_url("919876543210", "Namaste ji\n₹5")
        self.assertTrue(url.startswith("https://wa.me/919876543210?text="))
        self.assertNotIn(" ", url)
        self.assertEqual(whatsapp.chat_url(None, "x"), "https://wa.me/?text=x")
        m = whatsapp.reminder_message("Rahul", "SHREEJI BOXES", 150050, 63, "2026-08-01")
        self.assertIn("₹1,500.50", m)
        self.assertIn("01-08-2026", m)
        self.assertIn("advance", whatsapp.statement_message("R", "B", -500, "All dates"))
