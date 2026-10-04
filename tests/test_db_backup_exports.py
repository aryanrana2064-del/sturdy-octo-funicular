"""Database integrity, migrations, backup/restore, export edge cases."""
import os
import shutil
import sqlite3
import unittest
from pathlib import Path

from openpyxl import load_workbook

from tests.helpers import ServiceTestCase
from vx7khata import backup, db, exports
from vx7khata.service import KhataService


class TestDatabase(ServiceTestCase):
    def test_foreign_keys_enforced(self):
        with self.assertRaises(sqlite3.IntegrityError):
            self.svc.conn.execute(
                "INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at) "
                "VALUES (999, '2025-01-01', 'x', 100, 'UDHAAR', 'n', 'n')")

    def test_check_constraints(self):
        c = self.svc.add_customer("A")
        for amount, kind in ((0, "UDHAAR"), (-5, "JAMA"), (100, "OTHER")):
            with self.assertRaises(sqlite3.IntegrityError):
                self.svc.conn.execute(
                    "INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at) "
                    "VALUES (?, '2025-01-01', 'x', ?, ?, 'n', 'n')", (c.id, amount, kind))

    def test_transaction_rolls_back_on_error(self):
        with self.assertRaises(RuntimeError):
            with db.transaction(self.svc.conn):
                self.svc.conn.execute("INSERT INTO settings(key, value) VALUES('a', 'b')")
                raise RuntimeError("boom")
        self.assertEqual(self.svc.get_setting("a", "missing"), "missing")

    def test_settings_roundtrip(self):
        self.svc.set_setting("shop_name", "My Shop")
        self.svc.set_setting("shop_name", "Other Shop")
        self.assertEqual(self.svc.get_setting("shop_name"), "Other Shop")

    def test_schema_version_and_indexes(self):
        self.assertEqual(db.get_user_version(self.svc.conn), db.SCHEMA_VERSION)
        idx = {r[0] for r in self.svc.conn.execute("SELECT name FROM sqlite_master WHERE type='index'")}
        self.assertTrue({"idx_txn_customer_date", "idx_txn_date", "idx_customers_mobile", "idx_txn_item"} <= idx)

    def test_migration_from_v1_keeps_data_and_makes_safety_copy(self):
        path = Path(self.tmp.name) / "old.db"
        conn = sqlite3.connect(path, isolation_level=None)
        for stmt in db.MIGRATIONS[1]:
            conn.execute(stmt)
        conn.execute("PRAGMA user_version = 1")
        conn.execute("INSERT INTO customers(name, name_key, created_at, updated_at) VALUES('Old', 'old', 'n', 'n')")
        conn.execute("INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at) "
                     "VALUES (1, '2025-01-10', 'kept', 12300, 'UDHAAR', 'n', 'n')")
        conn.close()
        svc = KhataService(path)
        self.addCleanup(svc.close)
        self.assertEqual(db.get_user_version(svc.conn), db.SCHEMA_VERSION)
        self.assertEqual(svc.ledger(1).rows[0].txn.item, "kept")
        self.assertTrue((Path(self.tmp.name) / "old.db.pre-v1-migration.bak").exists())

    def test_rejects_newer_schema_and_foreign_files(self):
        newer = Path(self.tmp.name) / "newer.db"
        c = sqlite3.connect(newer)
        c.execute(f"PRAGMA user_version = {db.SCHEMA_VERSION + 1}")
        c.execute("CREATE TABLE x(a)")
        c.close()
        with self.assertRaises(db.DatabaseError):
            db.open_database(newer)
        other = Path(self.tmp.name) / "other.db"
        c = sqlite3.connect(other)
        c.execute("CREATE TABLE unrelated(a)")
        c.commit()
        c.close()
        with self.assertRaises(db.DatabaseError):
            db.open_database(other)
        junk = Path(self.tmp.name) / "junk.db"
        junk.write_bytes(b"this is not a database" * 50)
        with self.assertRaises(db.DatabaseError):
            db.open_database(junk)


class TestBackupRestore(ServiceTestCase):
    def setUp(self):
        super().setUp()
        self.c = self.svc.add_customer("Rahul", "9876543210")
        self.add(self.c.id, "UDHAAR", "10-01-2025", "Gold chain repair", "5000", notes="old")
        self.add(self.c.id, "JAMA", "15-01-2025", "Payment", "2000")
        self.svc.set_setting("shop_name", "My Shop")
        self.bk = Path(self.tmp.name) / "backup.db"

    def test_backup_is_valid_and_complete(self):
        self.svc.create_backup(self.bk)
        info = backup.validate_backup(self.bk)
        self.assertEqual((info.customers, info.transactions, info.schema_version), (1, 2, db.SCHEMA_VERSION))

    def test_restore_replaces_data_and_writes_safety_backup(self):
        self.svc.create_backup(self.bk)
        # change the live data after the backup was taken
        self.add(self.c.id, "UDHAAR", "01-03-2025", "later", "999")
        other = self.svc.add_customer("Amit")
        safety_dir = Path(self.tmp.name) / "safety"
        safety = self.svc.restore_backup(self.bk, safety_dir=safety_dir)

        # live service is usable immediately and shows the backed-up state
        names = [s.customer.name for s in self.svc.list_customers()]
        self.assertEqual(names, ["Rahul"])
        led = self.svc.ledger(self.c.id)
        self.assertEqual([r.txn.item for r in led.rows], ["Gold chain repair", "Payment"])
        self.assertEqual(led.closing_balance, 300000)
        self.assertEqual(led.rows[0].txn.txn_date, "2025-01-10")
        self.assertEqual(led.rows[0].txn.notes, "old")
        self.assertEqual(self.svc.get_setting("shop_name"), "My Shop")

        # the pre-restore state was kept and is itself a valid backup holding the later data
        self.assertTrue(safety.exists() and safety.parent == safety_dir)
        info = backup.validate_backup(safety)
        self.assertEqual((info.customers, info.transactions), (2, 3))
        self.assertFalse(Path(self.db_file + ".restore-tmp").exists())

    def test_restore_survives_reopen(self):
        self.svc.create_backup(self.bk)
        self.add(self.c.id, "UDHAAR", "01-03-2025", "later", "999")
        self.svc.restore_backup(self.bk, safety_dir=Path(self.tmp.name) / "s")
        self.svc.close()
        again = KhataService(self.db_file)
        self.addCleanup(again.close)
        self.assertEqual(again.ledger(self.c.id).closing_balance, 300000)

    def test_invalid_backups_rejected_and_live_data_untouched(self):
        self.svc.create_backup(self.bk)
        cases = {}
        cases["missing"] = Path(self.tmp.name) / "nope.db"
        empty = Path(self.tmp.name) / "empty.db"
        empty.write_bytes(b"")
        cases["empty"] = empty
        text = Path(self.tmp.name) / "text.db"
        text.write_text("hello")
        cases["text"] = text
        truncated = Path(self.tmp.name) / "trunc.db"
        truncated.write_bytes(self.bk.read_bytes()[:3000])
        cases["truncated"] = truncated
        corrupt = Path(self.tmp.name) / "corrupt.db"
        data = bytearray(self.bk.read_bytes())
        for i in range(4096, min(len(data), 4096 + 2000)):
            data[i] = 0xFF
        corrupt.write_bytes(bytes(data))
        cases["corrupt"] = corrupt
        foreign = Path(self.tmp.name) / "foreign.db"
        c = sqlite3.connect(foreign)
        c.execute("CREATE TABLE t(a)")
        c.commit()
        c.close()
        cases["foreign"] = foreign
        for name, path in cases.items():
            with self.subTest(name):
                with self.assertRaises(backup.BackupError):
                    self.svc.restore_backup(path, safety_dir=Path(self.tmp.name) / "s")
        self.assertEqual(self.svc.ledger(self.c.id).closing_balance, 300000)
        self.assertFalse((Path(self.tmp.name) / "s").exists() and any((Path(self.tmp.name) / "s").iterdir()))

    def test_backup_with_orphan_rows_rejected(self):
        bad = Path(self.tmp.name) / "orphan.db"
        shutil.copyfile(self.db_file, bad)
        c = sqlite3.connect(bad)
        c.execute("PRAGMA foreign_keys = OFF")
        c.execute("INSERT INTO transactions(customer_id, txn_date, item, amount_paise, txn_type, created_at, updated_at) "
                  "VALUES (77, '2025-01-01', 'orphan', 100, 'UDHAAR', 'n', 'n')")
        c.commit()
        c.close()
        with self.assertRaises(backup.BackupError):
            backup.validate_backup(bad)

    def test_cannot_backup_or_restore_onto_live_file(self):
        with self.assertRaises(backup.BackupError):
            self.svc.create_backup(self.db_file)
        with self.assertRaises(backup.BackupError):
            self.svc.restore_backup(self.db_file)

    def test_restore_older_schema_backup_gets_migrated(self):
        old = Path(self.tmp.name) / "v1.db"
        conn = sqlite3.connect(old, isolation_level=None)
        for stmt in db.MIGRATIONS[1]:
            conn.execute(stmt)
        conn.execute("PRAGMA user_version = 1")
        conn.execute("INSERT INTO customers(name, name_key, created_at, updated_at) VALUES('FromV1', 'fromv1', 'n', 'n')")
        conn.close()
        self.svc.restore_backup(old, safety_dir=Path(self.tmp.name) / "s")
        self.assertEqual(db.get_user_version(self.svc.conn), db.SCHEMA_VERSION)
        self.assertEqual([s.customer.name for s in self.svc.list_customers()], ["FromV1"])


class TestExportEdgeCases(ServiceTestCase):
    def test_special_characters_and_formula_text(self):
        c = self.svc.add_customer("Tom & Jerry <Traders>")
        self.add(c.id, "UDHAAR", "10-01-2025", "=1+1", "100", notes="a < b & c > d")
        self.add(c.id, "UDHAAR", "11-01-2025", "+SUM(A1)", "50")
        led = self.svc.ledger(c.id)
        xl = exports.export_statement_xlsx(led, Path(self.tmp.name) / "e.xlsx")
        ws = load_workbook(xl)["Statement"]
        values = [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]
        self.assertIn("=1+1", values)
        self.assertIn("+SUM(A1)", values)
        self.assertEqual(ws.cell(row=ws.max_row, column=2).value is not None, True)
        pdf = exports.export_statement_pdf(led, Path(self.tmp.name) / "e.pdf")
        self.assertTrue(pdf.exists() and pdf.stat().st_size > 1000)
        try:
            from pypdf import PdfReader
        except ImportError:
            return
        text = "\n".join(p.extract_text() for p in PdfReader(str(pdf)).pages)
        self.assertIn("Tom & Jerry <Traders>", text)
        self.assertIn("a < b & c > d", text)

    def test_many_rows_paginate_and_leave_no_part_files(self):
        c = self.svc.add_customer("Big")
        for i in range(150):
            self.add(c.id, "UDHAAR", "10-01-2025", f"item {i}", "10")
        led = self.svc.ledger(c.id)
        pdf = exports.export_statement_pdf(led, Path(self.tmp.name) / "big.pdf")
        try:
            from pypdf import PdfReader
        except ImportError:
            self.skipTest("pypdf not installed")
        reader = PdfReader(str(pdf))
        self.assertGreater(len(reader.pages), 2)
        text = "\n".join(p.extract_text() for p in reader.pages)
        self.assertIn("item 149", text)
        self.assertIn("1,500.00", text)
        self.assertEqual([p for p in os.listdir(self.tmp.name) if p.endswith(".part")], [])

    def test_register_and_customer_exports(self):
        a = self.svc.add_customer("A", "9876543210")
        b = self.svc.add_customer("B")
        self.add(a.id, "UDHAAR", "10-01-2025", "Gold ring", "1000")
        self.add(b.id, "JAMA", "11-01-2025", "Cash", "400")
        reg = self.svc.register()
        x = exports.export_register_xlsx(reg, Path(self.tmp.name) / "r.xlsx")
        ws = load_workbook(x)["Register"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        self.assertTrue(any(r[1] == "A" and r[2] == "Gold ring" and round(float(r[5]) * 100) == 100000 for r in rows))
        self.assertTrue(any(r[1] == "B" and r[2] == "Cash" and round(float(r[6]) * 100) == 40000 for r in rows))
        self.assertTrue(any(r[1] == "Net outstanding (udhaar - deposit)" and round(float(r[5]) * 100) == 60000 for r in rows))
        summaries = self.svc.list_customers()
        y = exports.export_customers_xlsx(summaries, Path(self.tmp.name) / "c.xlsx")
        ws = load_workbook(y).active
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        total = next(r for r in rows if r[1] == "Overall total")
        self.assertEqual([round(float(total[i]) * 100) for i in (4, 5, 6)], [100000, 40000, 60000])
        self.assertTrue(exports.export_register_pdf(reg, Path(self.tmp.name) / "r.pdf").exists())
        self.assertTrue(exports.export_customers_pdf(summaries, Path(self.tmp.name) / "c.pdf").exists())


if __name__ == "__main__":
    unittest.main()
