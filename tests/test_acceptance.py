"""Final acceptance tests A-F from the requirements."""
import os
import subprocess
import sys
import textwrap
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from tests.helpers import ServiceTestCase
from vx7khata import exports, money
from vx7khata.service import KhataService

ROOT = Path(__file__).resolve().parent.parent


class TestA_MultipleCustomers(ServiceTestCase):
    def test_rahul_balance_and_amit_separate(self):
        rahul = self.svc.add_customer("Rahul")
        amit = self.svc.add_customer("Amit")
        self.add(rahul.id, "UDHAAR", "01-02-2025", "Gold chain repair", "5000")
        self.add(rahul.id, "JAMA", "05-02-2025", "Previous payment", "1000")
        self.add(amit.id, "UDHAAR", "03-02-2025", "Personal udhaar", "2000")

        by_name = {s.customer.name: s for s in self.svc.list_customers()}
        self.assertEqual(by_name["Rahul"].total_udhaar, 500000)
        self.assertEqual(by_name["Rahul"].total_jama, 100000)
        self.assertEqual(by_name["Rahul"].balance, 400000)  # Rs 4,000
        self.assertEqual(by_name["Amit"].balance, 200000)  # untouched by Rahul's entries
        self.assertEqual(len(self.svc.ledger(rahul.id).rows), 2)
        self.assertEqual(len(self.svc.ledger(amit.id).rows), 1)
        self.assertEqual(self.svc.dashboard().net_outstanding, 600000)


class TestB_ItemCalculation(ServiceTestCase):
    def test_quantity_times_rate(self):
        c = self.svc.add_customer("Rahul")
        t = self.add(c.id, "UDHAAR", "01-02-2025", "Silver coin", quantity="3", rate="200")
        self.assertEqual(t.amount_paise, 60000)  # Rs 600
        self.assertEqual(money.format_inr(t.amount_paise), "₹600.00")
        self.assertEqual(t.quantity, "3")
        self.assertEqual(t.rate_paise, 20000)


class TestC_OldDateEntry(ServiceTestCase):
    def test_closing_balance_3000(self):
        c = self.svc.add_customer("Rahul")
        self.add(c.id, "UDHAAR", "10-01-2025", "Old udhaar", "5000")
        self.add(c.id, "JAMA", "15-01-2025", "Payment", "2000")
        led = self.svc.ledger(c.id)
        self.assertEqual(led.closing_balance, 300000)
        self.assertEqual([r.txn.txn_date for r in led.rows], ["2025-01-10", "2025-01-15"])

    def test_creation_timestamp_is_real_entry_time_not_transaction_date(self):
        c = self.svc.add_customer("Rahul")
        t = self.add(c.id, "UDHAAR", "10-01-2025", "Old udhaar", "5000")
        self.assertEqual(t.txn_date, "2025-01-10")
        self.assertTrue(t.created_at.startswith("2026-10-03"), t.created_at)


class TestD_HistoricalCorrection(ServiceTestCase):
    def test_insert_between_updates_balances(self):
        c = self.svc.add_customer("Rahul")
        self.add(c.id, "UDHAAR", "10-01-2025", "Old udhaar", "5000")
        self.add(c.id, "JAMA", "15-01-2025", "Payment", "2000")
        self.add(c.id, "UDHAAR", "12-01-2025", "Extra", "1000")
        led = self.svc.ledger(c.id)
        self.assertEqual(led.closing_balance, 400000)
        self.assertEqual(
            [(r.txn.txn_date, r.txn.item, r.balance) for r in led.rows],
            [("2025-01-10", "Old udhaar", 500000), ("2025-01-12", "Extra", 600000), ("2025-01-15", "Payment", 400000)],
        )


class TestE_Persistence(ServiceTestCase):
    def test_reopen_same_process(self):
        c = self.svc.add_customer("Rahul", "9876543210", "Main Road")
        self.add(c.id, "UDHAAR", "10-01-2025", "Gold chain repair", quantity="2", rate="750", notes="urgent")
        self.svc.close()
        again = KhataService(self.db_file, clock=self.clock)
        self.addCleanup(again.close)
        cust = again.list_customers()[0].customer
        self.assertEqual((cust.name, cust.mobile, cust.address), ("Rahul", "9876543210", "Main Road"))
        row = again.ledger(cust.id).rows[0].txn
        self.assertEqual((row.txn_date, row.item, row.quantity, row.amount_paise, row.notes),
                         ("2025-01-10", "Gold chain repair", "2", 150000, "urgent"))

    def test_reopen_in_a_separate_process(self):
        """Literally close the program and start a new one."""
        script_write = textwrap.dedent(f"""
            from vx7khata.service import KhataService
            s = KhataService({self.db_file!r})
            c = s.add_customer("Rahul")
            s.add_transaction(c.id, "UDHAAR", "10-01-2025", "Gold chain repair", amount="1500")
            s.close()
        """)
        script_read = textwrap.dedent(f"""
            from vx7khata.service import KhataService
            s = KhataService({self.db_file!r})
            c = s.list_customers()[0]
            led = s.ledger(c.customer.id)
            print(c.customer.name, led.rows[0].txn.txn_date, led.rows[0].txn.item, led.closing_balance)
        """)
        self.svc.close()
        env = dict(os.environ, PYTHONPATH=str(ROOT))
        subprocess.run([sys.executable, "-c", script_write], check=True, env=env, cwd=ROOT)
        out = subprocess.run([sys.executable, "-c", script_read], check=True, env=env, cwd=ROOT,
                             capture_output=True, text=True).stdout.strip()
        self.assertEqual(out, "Rahul 2025-01-10 Gold chain repair 150000")
        self.svc = KhataService(self.db_file, clock=self.clock)  # for cleanup


class TestF_Reports(ServiceTestCase):
    def setUp(self):
        super().setUp()
        self.c = self.svc.add_customer("Rahul", "9876543210", "Main Road")
        self.add(self.c.id, "UDHAAR", "10-01-2025", "Gold chain repair", "5000")
        self.add(self.c.id, "UDHAAR", "12-01-2025", "Silver coin", quantity="3", rate="200")
        self.add(self.c.id, "JAMA", "15-01-2025", "Cash payment", "2000", notes="received")
        self.out = Path(self.tmp.name)

    def test_excel_statement(self):
        led = self.svc.ledger(self.c.id)
        path = exports.export_statement_xlsx(led, self.out / "s.xlsx", shop_name="My Shop")
        ws = load_workbook(path)["Statement"]
        rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row)]
        # statement columns: date, item, udhaar, total deposit, deposit date, deposit balance, notes
        data = [r for r in rows if str(r[1]).split("\n")[0] in ("Gold chain repair", "Silver coin", "Cash payment")]
        self.assertEqual(len(data), 3)
        self.assertEqual([r[0].date() for r in data], [date(2025, 1, 10), date(2025, 1, 12), date(2025, 1, 15)])
        self.assertIn("Qty 3 × Rate ₹200.00", data[1][1])  # quantity x rate stays visible in the description
        self.assertEqual(round(float(data[1][2]) * 100), 60000)  # udhaar amount
        self.assertEqual(round(float(data[2][3]) * 100), 200000)  # deposit amount only
        self.assertEqual(data[2][4].date(), date(2025, 1, 15))  # deposit date
        self.assertEqual(data[2][6], "received")
        totals = next(r for r in rows if r[1] == "TOTAL")
        self.assertEqual(ws.cell(row=rows.index(totals) + 1, column=3).value[:5], "=SUM(")  # live formula
        self.assertTrue(any(r[1] == "Opening Balance" for r in rows))
        self.assertEqual(ws["B3"].value, "Rahul")
        values = load_workbook(path, data_only=True)["Statement"]
        vrows = [[cell.value for cell in row] for row in values.iter_rows()]
        vdata = [r for r in vrows if str(r[1]).split("\n")[0] in ("Gold chain repair", "Silver coin", "Cash payment")]
        self.assertEqual([round(float(r[5]) * 100) for r in vdata], [500000, 560000, 360000])
        vtotals = next(r for r in vrows if r[1] == "TOTAL")
        self.assertEqual([round(float(vtotals[i]) * 100) for i in (2, 3, 5)], [560000, 200000, 360000])

    def test_excel_filtered_period_has_correct_opening_balance(self):
        led = self.svc.ledger(self.c.id, date_from="12-01-2025", date_to="31-01-2025")
        ws = load_workbook(exports.export_statement_xlsx(led, self.out / "p.xlsx"))["Statement"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        opening = next(r for r in rows if r[1] == "Opening Balance")
        self.assertEqual(round(float(opening[5]) * 100), 500000)  # the 10-01 udhaar is before the range
        self.assertFalse(any(r[1] == "Gold chain repair" for r in rows))

    def test_pdf_statement(self):
        try:
            from pypdf import PdfReader
        except ImportError:
            self.skipTest("pypdf not installed")
        led = self.svc.ledger(self.c.id)
        path = exports.export_statement_pdf(led, self.out / "s.pdf", shop_name="My Shop")
        text = "\n".join(p.extract_text() for p in PdfReader(str(path)).pages)
        for needle in ("10-01-2025", "12-01-2025", "15-01-2025", "Gold chain repair", "Silver coin",
                       "Cash payment", "Rahul", "9876543210", "Opening Balance", "Final Deposit Balance",
                       "5,000.00", "5,600.00", "3,600.00", "2,000.00", "200.00", "600.00", "All dates"):
            self.assertIn(needle, text, needle)
        self.assertIn("₹", text)

    def test_item_filter_keeps_true_running_balance(self):
        led = self.svc.ledger(self.c.id, search="silver")
        self.assertEqual(len(led.rows), 1)
        self.assertEqual(led.rows[0].balance, 560000)  # includes the earlier 5,000 udhaar
        self.assertEqual(led.closing_balance, 360000)


if __name__ == "__main__":
    unittest.main()
