"""Import old khata data from an Excel (.xlsx) or CSV file.

Two layouts work (headers are matched by name, English or Hindi, in any order):
  * Customer | Mobile | Date | Item / Description | Udhaar | Jama | Qty | Rate | Notes | Deposit Date
  * Customer | Date | Item | Type (UDHAAR/JAMA) | Amount | ...
Everything is checked first (dry run). Nothing is saved unless every row is valid.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import from_excel

from .service import ImportReport, KhataService

MAX_ROWS = 20000

_ALIASES = {
    "customer": ["customer", "customer name", "name", "party", "party name", "naam", "नाम", "ग्राहक"],
    "mobile": ["mobile", "mobile no", "mobile number", "phone", "phone no", "contact", "whatsapp", "मोबाइल"],
    "date": ["date", "transaction date", "txn date", "tarikh", "तारीख", "दिनांक"],
    "item": ["item", "item description", "description", "details", "particulars", "vivaran", "माल", "विवरण"],
    "type": ["type", "txn type", "transaction type", "entry type"],
    "amount": ["amount", "total", "rakam", "रकम", "राशि"],
    "udhaar": ["udhaar", "udhar", "उधार", "debit", "due"],
    "jama": ["jama", "जमा", "credit", "deposit", "total deposit", "paid", "received"],
    "quantity": ["qty", "quantity", "nag", "मात्रा"],
    "rate": ["rate", "price", "bhav"],
    "notes": ["notes", "note", "remarks", "remark"],
    "deposit_date": ["deposit date", "jama date", "received date", "payment date"],
}
_LOOKUP = {alias: key for key, names in _ALIASES.items() for alias in names}
_UDHAAR_WORDS = {"udhaar", "udhar", "debit", "dr", "due", "उधार"}
_JAMA_WORDS = {"jama", "credit", "cr", "deposit", "received", "paid", "जमा"}


class ImportFileError(Exception):
    """The file cannot be read or does not have the needed columns."""


@dataclass
class ImportPlan:
    rows: list  # dicts ready for KhataService.import_rows
    file_errors: list  # [(line, message)] found while reading the file
    report: ImportReport  # dry-run result of the service checks

    @property
    def errors(self) -> list:
        return sorted(self.file_errors + list(self.report.errors), key=lambda e: e[0])

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.rows)


def _header_key(text) -> Optional[str]:
    cleaned = re.sub(r"[^\w\s]", " ", str(text or "").casefold())
    cleaned = re.sub(r"[_\s]+", " ", cleaned).strip()
    cleaned = re.sub(r"\b(rs|inr)\b", "", cleaned).strip()
    return _LOOKUP.get(cleaned)


def _num_text(value) -> str:
    """Spreadsheet number -> plain text without float noise (1200.5 -> '1200.5', 10.0 -> '10')."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        d = Decimal(repr(value))
        q = d.quantize(Decimal("0.01"))
        if abs(d - q) < Decimal("0.000001"):
            d = q
        text = format(d, "f")
        return text.rstrip("0").rstrip(".") if "." in text else text
    return str(value).strip()


def _cell_text(value, kind: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if kind in ("date", "deposit_date") and isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return from_excel(value).strftime("%d-%m-%Y")  # Excel serial number
        except (ValueError, OverflowError, TypeError):
            return str(value)
    if kind == "mobile" and isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(int(value))
    if isinstance(value, (int, float)):
        return _num_text(value)
    return str(value).strip()


def _read_table(path: Path) -> list[tuple[int, list]]:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # corrupt or not a workbook
            raise ImportFileError(f"This Excel file cannot be opened: {exc}") from None
        try:
            ws = wb.worksheets[0]
            return [(n, list(values)) for n, values in enumerate(ws.iter_rows(values_only=True), 1)]
        finally:
            wb.close()
    if ext in (".csv", ".txt"):
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="cp1252", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [(n, row) for n, row in enumerate(csv.reader(text.splitlines(), dialect), 1)]
    raise ImportFileError("Please choose an Excel (.xlsx) or CSV (.csv) file. Old .xls files: open and Save As .xlsx.")


def load_rows(path) -> tuple[list[dict], list]:
    """Read the file into service-ready row dicts. Returns (rows, errors found while reading)."""
    table = _read_table(Path(path))
    header_at = None
    for pos, (_n, cells) in enumerate(table[:15]):
        keys = [_header_key(c) for c in cells]
        if sum(1 for k in keys if k) >= 3 and "customer" in keys:
            header_at, columns = pos, {}
            for i, k in enumerate(keys):
                if k and k not in columns:
                    columns[k] = i
            break
    if header_at is None:
        raise ImportFileError("Could not find the header row. The first rows must have column names such as "
                              "Customer, Date, Item, Udhaar, Jama. Use 'Download template' for an example.")
    has_type = "type" in columns and "amount" in columns
    has_split = "udhaar" in columns or "jama" in columns
    if "date" not in columns:
        raise ImportFileError("The file needs a Date column.")
    if not (has_type or has_split):
        raise ImportFileError("The file needs either Udhaar and Jama columns, or a Type and an Amount column.")

    def get(cells, key):
        i = columns.get(key)
        return _cell_text(cells[i], key) if i is not None and i < len(cells) else ""

    rows: list[dict] = []
    errors: list = []
    for line, cells in table[header_at + 1:]:
        if not any(str(c).strip() for c in cells if c is not None):
            continue  # blank line
        if len(rows) + len(errors) >= MAX_ROWS:
            raise ImportFileError(f"The file has more than {MAX_ROWS} rows. Import it in smaller files.")
        udhaar, jama = get(cells, "udhaar"), get(cells, "jama")
        amount, kind = get(cells, "amount"), get(cells, "type").strip().casefold()
        if has_split and (udhaar or jama) and not has_type:
            if udhaar and jama:
                errors.append((line, "Both Udhaar and Jama are filled in this row. Use one row per entry."))
                continue
            txn_type, amount = ("UDHAAR", udhaar) if udhaar else ("JAMA", jama)
        elif has_type:
            if kind in _UDHAAR_WORDS:
                txn_type = "UDHAAR"
            elif kind in _JAMA_WORDS:
                txn_type = "JAMA"
            else:
                errors.append((line, f"Type '{get(cells, 'type')}' is not UDHAAR or JAMA"))
                continue
        else:
            errors.append((line, "Neither Udhaar nor Jama has an amount in this row"))
            continue
        rows.append({
            "line": line, "customer": get(cells, "customer"), "mobile": get(cells, "mobile"), "txn_type": txn_type,
            "txn_date": get(cells, "date"), "item": get(cells, "item"), "amount": amount,
            "quantity": get(cells, "quantity"), "rate": get(cells, "rate"), "notes": get(cells, "notes"),
            "deposit_date": get(cells, "deposit_date") or None,
        })
    return rows, errors


def plan_import(svc: KhataService, path) -> ImportPlan:
    """Read and fully check the file without saving anything."""
    rows, file_errors = load_rows(path)
    report = svc.import_rows(rows, dry_run=True)
    return ImportPlan(rows, file_errors, report)


def run_import(svc: KhataService, plan: ImportPlan) -> ImportReport:
    """Save a plan that has no errors. Raises ImportFileError otherwise."""
    if not plan.ok:
        raise ImportFileError("The file still has problems; nothing was imported.")
    return svc.import_rows(plan.rows)


def write_template(path) -> Path:
    """Blank import sheet with the right headers, plus an Example sheet and a short Help sheet."""
    path = Path(path)
    headers = ["Customer", "Mobile", "Date", "Item / Description", "Udhaar (₹)", "Jama (₹)", "Qty", "Rate (₹)",
               "Notes", "Deposit Date"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    fill = PatternFill("solid", fgColor="1F3A5F")
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = 20 if i != 4 else 34
    ws.freeze_panes = "A2"
    ex = wb.create_sheet("Example")
    ex.append(headers)
    ex.append(["Rahul Sharma", "9876543210", "10-01-2025", "Gold box set", 5000, None, None, None, "", None])
    ex.append(["Rahul Sharma", "", "12-01-2025", "Cartons", 3000, None, 10, 300, "page 4", None])
    ex.append(["Rahul Sharma", "", "20-01-2025", "Cash received", None, 2000, None, None, "", "20-01-2025"])
    for i in range(1, len(headers) + 1):
        ex.cell(row=1, column=i).font = Font(bold=True)
        ex.column_dimensions[ex.cell(row=1, column=i).column_letter].width = 20
    help_ws = wb.create_sheet("Help")
    for line in (
        "Fill the 'Data' sheet, one row per entry. Do not rename the column headings.",
        "Customer + Date + Item + (Udhaar OR Jama) are needed. Jama rows need no Item.",
        "Dates: DD-MM-YYYY (old dates are fine). Amounts: plain numbers, up to 2 decimals.",
        "Qty and Rate are optional notes about the item; always fill the Udhaar amount yourself.",
        "A new customer name is created automatically; a name already in the app is matched (case ignored).",
        "Nothing is imported unless every row is valid. The app shows every problem with its row number.",
    ):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    wb.save(path)
    return path
